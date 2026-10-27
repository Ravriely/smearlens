from dataclasses import dataclass, field
from functools import cache
from pathlib import Path
from typing import Iterator

from PIL import Image

CLASSES = ("WBC", "RBC", "Platelet")
ALLOWED_SOURCES = frozenset({"PBC", "BCCD", "BCDD"})
SPLITS = ("train", "val", "test")


@dataclass(frozen=True)
class Box:
    label: str
    x0: float
    y0: float
    x1: float
    y1: float

    @property
    def area(self) -> float:
        return max(0.0, self.x1 - self.x0) * max(0.0, self.y1 - self.y0)


@dataclass
class Sample:
    image_path: Path
    source: str
    width: int
    height: int
    boxes: list[Box] = field(default_factory=list)

    def count(self, label: str) -> int:
        return sum(1 for b in self.boxes if b.label == label)


def parse_yolo_line(line: str, width: int, height: int) -> Box | None:
    parts = line.split()
    if len(parts) != 5:
        return None
    cls, cx, cy, w, h = int(parts[0]), *map(float, parts[1:])
    return Box(
        CLASSES[cls],
        max(0.0, (cx - w / 2) * width),
        max(0.0, (cy - h / 2) * height),
        min(width, (cx + w / 2) * width),
        min(height, (cy + h / 2) * height),
    )


@cache
def load_sources(root: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    sheet = load_workbook(root / "metadata_file.xlsx", read_only=True).active
    rows = sheet.iter_rows(min_row=2, values_only=True)
    return {row[2]: row[1] for row in rows if row[2]}


def txl_pbc_root(data_dir: Path) -> Path:
    return data_dir / "txl-pbc" / "TXL-PBC"


def iter_txl_pbc(data_dir: Path, split: str) -> Iterator[Sample]:
    root = txl_pbc_root(data_dir)
    sources = load_sources(root)
    for image_path in sorted((root / "images" / split).iterdir()):
        source = sources.get(image_path.name)
        if source not in ALLOWED_SOURCES:
            continue
        with Image.open(image_path) as image:
            width, height = image.size
        sample = Sample(image_path, source, width, height)
        label_path = root / "labels" / split / f"{image_path.stem}.txt"
        if label_path.exists():
            for line in label_path.read_text().splitlines():
                box = parse_yolo_line(line, width, height)
                if box:
                    sample.boxes.append(box)
        yield sample
