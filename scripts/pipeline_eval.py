import argparse
import json
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
from PIL import Image, ImageDraw

from smearlens.datasets import Box, iter_txl_pbc, txl_pbc_root
from smearlens.metrics import MatchResult, contains_center, match_centers
from smearlens.muse import ask_image, parse_label
from smearlens.pipeline import (
    CANDIDATE_CLASSES,
    NUCLEATED,
    PROMPT,
    Candidate,
    crop,
    crop_at_scale,
    dedupe,
    differential,
    pbc_label,
    probe_labels,
    rbc_scale,
    review_flags,
)

COLORS = {
    "neutrophil": (37, 99, 235),
    "lymphocyte": (22, 163, 74),
    "monocyte": (217, 119, 6),
    "eosinophil": (220, 38, 38),
    "basophil": (124, 58, 237),
    "ig": (219, 39, 119),
    "erythroblast": (8, 145, 178),
}


def original_names(data_dir: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    sheet = load_workbook(txl_pbc_root(data_dir) / "metadata_file.xlsx", read_only=True).active
    return {row[2]: row[0] for row in sheet.iter_rows(min_row=2, values_only=True) if row[2]}


def load_cache(path: Path) -> dict[str, dict]:
    if not path.exists():
        return {}
    return {entry["key"]: entry for entry in map(json.loads, path.read_text().splitlines())}


def cache_key(image: str, box: Box) -> str:
    return f"{image}:{box.x0:.0f},{box.y0:.0f},{box.x1:.0f},{box.y1:.0f}"


def glimmer_features(keys: list[str], crops: list[bytes], path: Path) -> dict[str, np.ndarray]:
    stored = dict(np.load(path)) if path.exists() else {}
    missing = [(k, c) for k, c in zip(keys, crops) if k not in stored]
    if missing:
        import modal

        from modal_glimmer_probe import app, embed_crops

        with modal.enable_output(), app.run():
            features = embed_crops.remote([c for _, c in missing])
        stored.update({k: f for (k, _), f in zip(missing, features)})
        np.savez(path, **stored)
    return {k: stored[k] for k in keys}


def draw(image: Image.Image, cells: list[Candidate], truth: list[Box]) -> Image.Image:
    canvas = image.convert("RGB")
    pen = ImageDraw.Draw(canvas)
    for b in truth:
        pen.rectangle((b.x0, b.y0, b.x1, b.y1), outline=(255, 255, 255), width=1)
    for c in cells:
        color = COLORS.get(c.label, (0, 0, 0))
        pen.rectangle((c.box.x0, c.box.y0, c.box.x1, c.box.y1), outline=color, width=3)
        pen.text((c.box.x0 + 3, c.box.y0 + 2), c.label, fill=color)
    return canvas


def main() -> None:
    parser = argparse.ArgumentParser(description="End-to-end SmearLens: SAM 3 candidates, Muse Spark labels.")
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument("--detections", type=Path, default=Path("runs/pipeline-detect/detections.json"))
    parser.add_argument("--split", default="test")
    parser.add_argument("--images", type=int, default=10)
    parser.add_argument("--min-score", type=float, default=0.3)
    parser.add_argument("--max-tokens", type=int, default=4000)
    parser.add_argument("--workers", type=int, default=6)
    parser.add_argument("--max-calls", type=int, default=700)
    parser.add_argument("--typer", choices=["muse", "glimmer"], default="muse")
    parser.add_argument("--probe", type=Path, default=Path("runs/glimmer/probe.npz"))
    parser.add_argument("--glimmer-crop", choices=["tight", "rbc"], default="rbc")
    parser.add_argument("--recheck-platelets", action="store_true")
    parser.add_argument("--out", type=Path, default=Path("results/pipeline"))
    args = parser.parse_args()
    suffix = "" if args.typer == "muse" else f"_{args.typer}_{args.glimmer_crop}" + ("_plt" if args.recheck_platelets else "")

    raw = json.loads(args.detections.read_text())
    detections = {d["image"]: d for d in raw if d["label"] == "WBC"}
    rbc_px = {
        d["image"]: rbc_scale([max(b[2] - b[0], b[3] - b[1]) for b, s in zip(d["boxes"], d["scores"]) if s >= 0.5])
        for d in raw
        if d["label"] == "RBC"
    }
    samples = [s for s in iter_txl_pbc(args.data_dir, args.split) if s.image_path.name in detections]
    samples = samples[: args.images]
    originals = original_names(args.data_dir)

    candidates: dict[str, list[Candidate]] = {}
    for sample in samples:
        d = detections[sample.image_path.name]
        found = [
            Candidate(Box("WBC", *box), score)
            for box, score in zip(d["boxes"], d["scores"])
            if score >= args.min_score
        ]
        candidates[sample.image_path.name] = dedupe(found)

    args.out.mkdir(parents=True, exist_ok=True)
    cache_path = args.out / "labels.jsonl"
    cache = load_cache(cache_path)
    todo = [
        (sample, c)
        for sample in samples
        for c in candidates[sample.image_path.name]
        if cache_key(sample.image_path.name, c.box) not in cache
    ]
    if len(todo) > args.max_calls:
        raise SystemExit(f"{len(todo)} calls planned, above --max-calls {args.max_calls}")
    print(f"{sum(len(v) for v in candidates.values())} candidates, {len(todo)} new calls", flush=True)

    images = {s.image_path.name: Image.open(s.image_path).convert("RGB") for s in samples}

    def label(job):
        sample, c = job
        text, tokens = ask_image(crop(images[sample.image_path.name], c.box), PROMPT, args.max_tokens)
        return {
            "key": cache_key(sample.image_path.name, c.box),
            "label": parse_label(text, CANDIDATE_CLASSES),
            "text": text.replace("\n", " "),
            "tokens": tokens,
        }

    with cache_path.open("a") as f, ThreadPoolExecutor(args.workers) as pool:
        for i, entry in enumerate(pool.map(label, todo), 1):
            cache[entry["key"]] = entry
            f.write(json.dumps(entry) + "\n")
            f.flush()
            if i % 25 == 0:
                print(f"  {i}/{len(todo)}", flush=True)

    for sample in samples:
        name = sample.image_path.name
        for c in candidates[name]:
            entry = cache[cache_key(name, c.box)]
            c.label, c.text = entry["label"], entry["text"]

    muse_labels: dict[str, str] = {}
    if args.typer == "glimmer":
        typed = [
            (s, c)
            for s in samples
            for c in candidates[s.image_path.name]
            if c.label in NUCLEATED or (args.recheck_platelets and c.label == "platelet")
        ]
        keys = [cache_key(s.image_path.name, c.box) for s, c in typed]
        if args.glimmer_crop == "rbc":
            crops = [crop_at_scale(images[s.image_path.name], c.box, rbc_px[s.image_path.name]) for s, c in typed]
        else:
            crops = [crop(images[s.image_path.name], c.box) for s, c in typed]
        features = glimmer_features(keys, crops, args.out / f"glimmer_features_{args.glimmer_crop}.npz")
        labels = probe_labels([features[k] for k in keys], dict(np.load(args.probe)))
        for (s, c), key, new in zip(typed, keys, labels):
            muse_labels[key] = c.label
            c.label = new

    before, after = MatchResult(), MatchResult()
    errors_before, errors_after = [], []
    subtype = Counter()
    confusion: dict[str, Counter] = defaultdict(Counter)
    labels_seen = Counter()
    reports = []
    for sample in samples:
        name = sample.image_path.name
        for c in candidates[name]:
            labels_seen[c.label] += 1
        cells = [c for c in candidates[name] if c.label in NUCLEATED]
        truth = [b for b in sample.boxes if b.label == "WBC"]
        before += match_centers([c.box for c in candidates[name]], truth)
        after += match_centers([c.box for c in cells], truth)
        errors_before.append(abs(len(candidates[name]) - len(truth)))
        errors_after.append(abs(len(cells) - len(truth)))

        known = pbc_label(originals.get(name, "")) if sample.source == "PBC" else None
        if known:
            for t in truth:
                hits = [c for c in cells if contains_center(t, c.box)]
                pred = max(hits, key=lambda c: c.score).label if hits else "missed"
                subtype["correct" if pred == known else "wrong"] += 1
                confusion[known][pred] += 1

        reports.append(
            {
                "image": name,
                "source": sample.source,
                "known_wbc_type": known,
                "true_wbc": len(truth),
                "found_wbc": len(cells),
                "differential": {k: round(v, 3) for k, v in differential(cells).items() if v},
                "review": review_flags(candidates[name]),
            }
        )
        overlays = args.out / f"overlays{suffix}"
        overlays.mkdir(exist_ok=True)
        draw(images[name], cells, truth).save(overlays / f"{sample.image_path.stem}.jpg")

    total_subtype = subtype["correct"] + subtype["wrong"]
    tokens = sum(cache[cache_key(s.image_path.name, c.box)]["tokens"] for s in samples for c in candidates[s.image_path.name])
    agreement = sum(
        muse_labels[cache_key(s.image_path.name, c.box)] == c.label
        for s in samples
        for c in candidates[s.image_path.name]
        if cache_key(s.image_path.name, c.box) in muse_labels
    )
    summary = {
        "typer": args.typer,
        "muse_probe_agreement": f"{agreement}/{len(muse_labels)}" if muse_labels else None,
        "images": len(samples),
        "candidates": sum(len(v) for v in candidates.values()),
        "min_score": args.min_score,
        "candidate_labels": dict(labels_seen.most_common()),
        "wbc_detection": {
            "sam_only": {"precision": round(before.precision, 3), "recall": round(before.recall, 3), "f1": round(before.f1, 3)},
            "sam_plus_muse": {"precision": round(after.precision, 3), "recall": round(after.recall, 3), "f1": round(after.f1, 3)},
        },
        "wbc_count_mean_abs_error": {
            "sam_only": round(float(np.mean(errors_before)), 2),
            "sam_plus_muse": round(float(np.mean(errors_after)), 2),
        },
        "pbc_subtype": {
            "cells": total_subtype,
            "accuracy": round(subtype["correct"] / total_subtype, 3) if total_subtype else None,
            "confusion": {k: dict(v) for k, v in confusion.items()},
        },
        "tokens": tokens,
    }
    (args.out / f"summary{suffix}.json").write_text(json.dumps(summary, indent=2))
    (args.out / f"reports{suffix}.json").write_text(json.dumps(reports, indent=2))
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
