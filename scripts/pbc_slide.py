"""Assemble a synthetic slide from the TXL-PBC fields of one source and score the backend on it.

    uv run python scripts/pbc_slide.py fields --out runs/slide-pbc/truth.json
    modal run scripts/modal_app.py $(uv run python scripts/pbc_slide.py args runs/slide-pbc/truth.json) \
        --out runs/slide-pbc/reports.json
    uv run python scripts/pbc_slide.py score runs/slide-pbc/truth.json runs/slide-pbc/reports.json
"""

import argparse
import json
from collections import Counter
from pathlib import Path

from smearlens.datasets import Box, iter_txl_pbc, txl_pbc_root
from smearlens.metrics import contains_center
from smearlens.pipeline import NUCLEATED, pbc_label


def build(data_dir: Path, source: str, splits: list[str]) -> list[dict]:
    from openpyxl import load_workbook

    sheet = load_workbook(txl_pbc_root(data_dir) / "metadata_file.xlsx", read_only=True).active
    original = {r[2]: r[0] for r in sheet.iter_rows(min_row=2, values_only=True) if r[2]}
    return [
        {
            "path": str(s.image_path),
            "split": split,
            "original": original.get(s.image_path.name),
            "type": pbc_label(original.get(s.image_path.name) or ""),
            "wbc_boxes": [[b.x0, b.y0, b.x1, b.y1] for b in s.boxes if b.label == "WBC"],
            "red_cells": s.count("RBC"),
            "platelets": s.count("Platelet"),
        }
        for split in splits
        for s in iter_txl_pbc(data_dir, split)
        if s.source == source
    ]


def central(boxes: list[list[float]], width: float, height: float) -> list[float] | None:
    def distance(b):
        return ((b[0] + b[2]) / 2 - width / 2) ** 2 + ((b[1] + b[3]) / 2 - height / 2) ** 2

    return min(boxes, key=distance) if boxes else None


def percent(counts: Counter) -> dict[str, float]:
    total = sum(counts.values())
    return {t: round(100 * counts[t] / total, 1) for t in NUCLEATED} if total else {}


def score(truth: list[dict], reports: list[dict]) -> dict:
    by_path = {r["path"]: r for r in reports}
    rows = {}
    for split in sorted({t["split"] for t in truth}) + ["all"]:
        part = [t for t in truth if split in ("all", t["split"])]
        typed = right = missed = review = review_wrong = 0
        wbc_true = wbc_found = matched = 0
        predicted, expected = Counter(), Counter()
        count_error = {"red_cells": 0.0, "platelets": 0.0}
        for t in part:
            r = by_path[t["path"]]
            cells = r["white_cells"]
            truth_boxes = [Box("WBC", *b) for b in t["wbc_boxes"]]
            wbc_true += len(truth_boxes)
            wbc_found += len(cells)
            matched += sum(any(contains_center(tb, Box("WBC", *c["box"])) for tb in truth_boxes) for c in cells)
            predicted.update(c["type"] for c in cells)
            count_error["red_cells"] += abs(r["red_cell_count"] - t["red_cells"])
            count_error["platelets"] += abs(r["platelet_count"] - t["platelets"])
            if not t["type"]:
                continue
            expected[t["type"]] += 1
            target = central(t["wbc_boxes"], r["image"]["width"], r["image"]["height"])
            hit = next((c for c in cells if target and contains_center(Box("WBC", *target), Box("WBC", *c["box"]))), None)
            if not hit:
                missed += 1
                continue
            typed += 1
            right += hit["type"] == t["type"]
            review += hit["needs_review"]
            review_wrong += hit["needs_review"] and hit["type"] != t["type"]
        wrong = typed - right
        rows[split] = {
            "fields": len(part),
            "white_cells_truth": wbc_true,
            "white_cells_found": wbc_found,
            "found_on_a_truth_box": matched,
            "labelled_cells": sum(expected.values()),
            "labelled_cells_missed": missed,
            "type_accuracy": round(right / typed, 4) if typed else None,
            "wrong_type": wrong,
            "wrong_type_marked_for_review": review_wrong,
            "labelled_marked_for_review": review,
            "differential_predicted": percent(predicted),
            "differential_truth": percent(expected),
            "count_error_per_field": {k: round(v / len(part), 2) for k, v in count_error.items()},
        }
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest="command", required=True)
    fields = sub.add_parser("fields")
    fields.add_argument("--data-dir", type=Path, default=Path("data"))
    fields.add_argument("--source", default="PBC")
    fields.add_argument("--splits", nargs="+", default=["test", "val"])
    fields.add_argument("--out", type=Path, required=True)
    args_cmd = sub.add_parser("args")
    args_cmd.add_argument("truth", type=Path)
    score_cmd = sub.add_parser("score")
    score_cmd.add_argument("truth", type=Path)
    score_cmd.add_argument("reports", type=Path)
    score_cmd.add_argument("--out", type=Path)
    args = parser.parse_args()

    if args.command == "fields":
        rows = build(args.data_dir, args.source, args.splits)
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(rows, indent=1))
        for split in args.splits:
            part = [r for r in rows if r["split"] == split]
            print(split, len(part), "fields,", sum(len(r["wbc_boxes"]) for r in part), "white cells,",
                  dict(Counter(r["type"] for r in part)))
    elif args.command == "args":
        print(" ".join(f"--image {r['path']}" for r in json.loads(args.truth.read_text())))
    else:
        result = score(json.loads(args.truth.read_text()), json.loads(args.reports.read_text()))
        text = json.dumps(result, indent=2)
        if args.out:
            args.out.write_text(text + "\n")
        print(text)


if __name__ == "__main__":
    main()
