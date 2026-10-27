import csv
import hashlib
from dataclasses import dataclass
from pathlib import Path

from smearlens.datasets import load_sources, txl_pbc_root

PBC_CLASSES = (
    "basophil",
    "eosinophil",
    "erythroblast",
    "ig",
    "lymphocyte",
    "monocyte",
    "neutrophil",
    "platelet",
)


@dataclass(frozen=True)
class Cell:
    path: Path
    label: str
    split: str


def pbc_root(data_dir: Path) -> Path:
    return data_dir / "pbc" / "PBC_dataset_normal_DIB"


def hash_split(name: str, val: float = 0.15, test: float = 0.15) -> str:
    bucket = int(hashlib.sha256(name.encode()).hexdigest()[:8], 16) / 0xFFFFFFFF
    if bucket < test:
        return "test"
    if bucket < test + val:
        return "val"
    return "train"


def txl_pbc_splits(data_dir: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    root = txl_pbc_root(data_dir)
    sources = load_sources(root)
    new_to_split = {
        p.name: split
        for split in ("train", "val", "test")
        for p in (root / "images" / split).iterdir()
    }
    sheet = load_workbook(root / "metadata_file.xlsx", read_only=True).active
    return {
        row[0]: new_to_split[row[2]]
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[2] in new_to_split and sources.get(row[2]) == "PBC"
    }


def build_splits(data_dir: Path) -> list[Cell]:
    pinned = txl_pbc_splits(data_dir)
    cells = []
    for label in PBC_CLASSES:
        for path in sorted((pbc_root(data_dir) / label).glob("*.jpg")):
            split = pinned.get(path.name) or hash_split(path.name)
            cells.append(Cell(path, label, split))
    return cells


def write_splits(cells: list[Cell], out: Path, data_dir: Path) -> None:
    with out.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["path", "label", "split"])
        for cell in cells:
            writer.writerow([cell.path.relative_to(data_dir).as_posix(), cell.label, cell.split])


def read_splits(path: Path, data_dir: Path) -> list[Cell]:
    with path.open() as f:
        return [Cell(data_dir / row["path"], row["label"], row["split"]) for row in csv.DictReader(f)]
